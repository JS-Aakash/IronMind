import hashlib
import logging
from datetime import datetime
from pathlib import Path
from typing import Optional

from packages.shared.utils.helpers import generate_uuid
from services.artifacts.generators.base import BaseArtifactGenerator
from services.artifacts.models import GeneratedArtifactRecord, SpreadsheetData

logger = logging.getLogger(__name__)


class XlsxSpreadsheetGenerator(BaseArtifactGenerator):
    """Generates real Microsoft Excel (.xlsx) workbooks using openpyxl."""

    @property
    def artifact_type(self) -> str:
        return "xlsx"

    async def generate(
        self,
        data: SpreadsheetData,
        output_dir: Path,
        task_id: Optional[str] = None,
    ) -> GeneratedArtifactRecord:
        import openpyxl
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        # Style definitions
        header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        data_font = Font(name="Calibri", size=10)
        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1"),
        )

        for sheet_name, sheet_content in data.sheets.items():
            ws = wb.create_sheet(title=sheet_name[:31])
            headers = sheet_content.get("headers", [])
            rows = sheet_content.get("rows", [])

            # Write headers
            if headers:
                ws.append(headers)
                for col_idx in range(1, len(headers) + 1):
                    cell = ws.cell(row=1, column=col_idx)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.border = thin_border

            # Write rows
            for row_idx, row_data in enumerate(rows, start=2):
                ws.append(row_data)
                for col_idx in range(1, len(row_data) + 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.font = data_font
                    cell.border = thin_border
                    cell.alignment = Alignment(horizontal="left", vertical="center")

            # Auto-adjust column widths
            for col in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
        clean_title = data.title.replace(" ", "_").replace("-", "_")
        filename = f"{clean_title}_{timestamp}.xlsx"
        file_path = output_dir / filename
        wb.save(str(file_path))

        file_bytes = file_path.read_bytes()
        sha256 = hashlib.sha256(file_bytes).hexdigest()
        artifact_id = generate_uuid("ART")

        return GeneratedArtifactRecord(
            artifact_id=artifact_id,
            task_id=task_id or data.task_id,
            filename=filename,
            type="xlsx",
            created_at=datetime.utcnow().isoformat(),
            source_documents=data.source_documents,
            models_used=data.models_used,
            verification_status="verified",
            sha256_hash=sha256,
            size_bytes=len(file_bytes),
            file_path=str(file_path),
            download_url=f"/api/v1/artifacts/{artifact_id}/download",
            metadata={"title": data.title, "sheets_count": len(data.sheets)},
        )
