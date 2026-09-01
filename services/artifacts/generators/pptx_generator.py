import hashlib
import logging
from datetime import datetime
from pathlib import Path
from typing import Optional

from packages.shared.utils.helpers import generate_uuid
from services.artifacts.generators.base import BaseArtifactGenerator
from services.artifacts.models import GeneratedArtifactRecord, PresentationData

logger = logging.getLogger(__name__)


class PptxPresentationGenerator(BaseArtifactGenerator):
    """Generates real Microsoft PowerPoint (.pptx) presentations using python-pptx."""

    @property
    def artifact_type(self) -> str:
        return "pptx"

    async def generate(
        self,
        data: PresentationData,
        output_dir: Path,
        task_id: Optional[str] = None,
    ) -> GeneratedArtifactRecord:
        import pptx
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor

        prs = pptx.Presentation()
        prs.slide_width = Inches(13.33)
        prs.slide_height = Inches(7.5)

        blank_slide_layout = prs.slide_layouts[6]

        # 1. Title Slide
        title_slide = prs.slides.add_slide(blank_slide_layout)
        tx_box = title_slide.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.33), Inches(3.5))
        tf = tx_box.text_frame
        tf.word_wrap = True

        p_main = tf.paragraphs[0]
        p_main.text = data.title
        p_main.font.bold = True
        p_main.font.size = Pt(40)
        p_main.font.color.rgb = RGBColor(15, 23, 42)

        if data.subtitle:
            p_sub = tf.add_paragraph()
            p_sub.text = data.subtitle
            p_sub.font.size = Pt(20)
            p_sub.font.color.rgb = RGBColor(14, 165, 233)

        p_meta = tf.add_paragraph()
        p_meta.text = f"\nMANGALORE REFINERY AND PETROCHEMICALS LIMITED | {datetime.utcnow().strftime('%B %Y')}"
        p_meta.font.size = Pt(12)
        p_meta.font.color.rgb = RGBColor(100, 116, 139)

        # 2. Content Slides
        for s_idx, slide_data in enumerate(data.slides, 1):
            slide = prs.slides.add_slide(blank_slide_layout)
            
            # Slide Header
            h_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(1.0))
            h_tf = h_box.text_frame
            h_p = h_tf.paragraphs[0]
            h_p.text = f"{s_idx}. {slide_data.get('title', 'Slide')}"
            h_p.font.bold = True
            h_p.font.size = Pt(24)
            h_p.font.color.rgb = RGBColor(15, 23, 42)

            # Bullet points
            c_box = slide.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.33), Inches(4.8))
            c_tf = c_box.text_frame
            c_tf.word_wrap = True

            bullets = slide_data.get("bullet_points", [])
            for b_idx, bullet in enumerate(bullets):
                p = c_tf.paragraphs[0] if b_idx == 0 else c_tf.add_paragraph()
                p.text = f"•  {bullet}"
                p.font.size = Pt(16)
                p.font.color.rgb = RGBColor(51, 65, 85)
                p.space_after = Pt(14)

        timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
        clean_title = data.title.replace(" ", "_").replace("-", "_")
        filename = f"{clean_title}_{timestamp}.pptx"
        file_path = output_dir / filename
        prs.save(str(file_path))

        file_bytes = file_path.read_bytes()
        sha256 = hashlib.sha256(file_bytes).hexdigest()
        artifact_id = generate_uuid("ART")

        return GeneratedArtifactRecord(
            artifact_id=artifact_id,
            task_id=task_id or data.task_id,
            filename=filename,
            type="pptx",
            created_at=datetime.utcnow().isoformat(),
            source_documents=data.source_documents,
            models_used=data.models_used,
            verification_status="verified",
            sha256_hash=sha256,
            size_bytes=len(file_bytes),
            file_path=str(file_path),
            download_url=f"/api/v1/artifacts/{artifact_id}/download",
            metadata={"title": data.title, "slides_count": len(data.slides) + 1},
        )
